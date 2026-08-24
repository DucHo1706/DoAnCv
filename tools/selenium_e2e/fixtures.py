from __future__ import annotations

import json
import re
import csv
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import textwrap
from typing import Any


FICTIONAL_NAMES = [
    "Nguyễn Minh An", "Trần Gia Bảo", "Lê Hoàng Duy", "Phạm Khánh Linh", "Võ Quốc Huy",
    "Đặng Thu Hà", "Bùi Đức Long", "Đỗ Ngọc Mai", "Hồ Anh Khoa", "Ngô Thanh Trúc",
    "Dương Nhật Nam", "Lý Quỳnh Anh", "Mai Tuấn Kiệt", "Tạ Mỹ Duyên", "Trịnh Hải Đăng",
    "Nguyễn Thảo Vy", "Trần Minh Quân", "Lê Bảo Ngọc", "Phạm Hữu Phước", "Võ Kim Chi",
]

MATCH_SCENARIOS = (
    "strong_same_role",
    "partial_same_domain",
    "negative_cross_domain",
)

EVIDENCE_CASE_TYPES = (
    "strong_documented",
    "preferred_skill_gap",
    "experience_below_requirement",
    "transferable_education",
    "language_needs_verification",
)

# Xen kẽ định dạng ngay từ đầu batch. Nếu lượt chạy bị gián đoạn, bằng chứng vẫn
# bao phủ DOCX, PDF text và PDF scan thay vì chỉ mới đi qua toàn bộ DOCX.
LAYOUT_VARIANT_ORDER = (
    0, 9, 12, 15, 1, 10, 13, 16, 2, 11,
    14, 17, 3, 18, 4, 19, 5, 6, 7, 8,
)


@dataclass(frozen=True)
class CandidateFixture:
    index: int
    source_cv_id: str
    full_name: str
    email: str
    case_type: str
    match_scenario: str
    expected_offline_score: int
    document_path: Path
    layout: str


def _load_json(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, list):
        raise ValueError(f"Dữ liệu {path.name} phải là một mảng JSON.")
    return payload


def _load_expected_scores(path: Path) -> dict[tuple[str, str, str], int]:
    if not path.exists():
        raise FileNotFoundError(
            f"Thiếu {path}. Hãy chạy Python/tools/run_offline_algorithm_benchmark.py trước."
        )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return {
            (row["cv_id"], row["job_id"], row["scenario"]): int(row["score"])
            for row in csv.DictReader(handle)
        }


def _select_stratified_pairs(
    cvs: list[dict[str, Any]],
    pairs: list[dict[str, Any]],
    benchmark_job_id: str,
    count: int,
) -> list[tuple[dict[str, Any], str]]:
    """Chọn theo ma trận quan hệ CV-JD x chất lượng bằng chứng, không thiên về ca mạnh."""
    cv_by_id = {str(cv["id"]): cv for cv in cvs}
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for pair in pairs:
        if pair.get("job_id") != benchmark_job_id:
            continue
        cv = cv_by_id.get(str(pair.get("cv_id")))
        scenario = str(pair.get("scenario", ""))
        if cv is None or scenario not in MATCH_SCENARIOS:
            continue
        grouped[(scenario, str(cv.get("case_type", "unknown")))].append(cv)

    ordered: list[tuple[dict[str, Any], str]] = []
    seen: set[str] = set()

    # Vòng đầu tạo đủ lưới 3 x 5 khi count=15: mỗi mức phù hợp có đủ năm
    # tình huống bằng chứng. Thứ tự xen kẽ giúp cả batch nhỏ cũng không toàn 100%.
    for case_type in EVIDENCE_CASE_TYPES:
        for scenario in MATCH_SCENARIOS:
            candidates = grouped.get((scenario, case_type), [])
            candidate = next(
                (item for item in candidates if str(item["id"]) not in seen),
                None,
            )
            if candidate is None:
                continue
            ordered.append((candidate, scenario))
            seen.add(str(candidate["id"]))

    # Nếu cần hơn 15 CV, lấy phần còn lại theo round-robin giữa ba dải điểm.
    remaining_by_scenario: dict[str, list[dict[str, Any]]] = {
        scenario: [
            cv_by_id[str(pair["cv_id"])]
            for pair in pairs
            if pair.get("job_id") == benchmark_job_id
            and pair.get("scenario") == scenario
            and str(pair.get("cv_id")) in cv_by_id
            and str(pair.get("cv_id")) not in seen
        ]
        for scenario in MATCH_SCENARIOS
    }
    while len(ordered) < count and any(remaining_by_scenario.values()):
        for scenario in MATCH_SCENARIOS:
            remaining = remaining_by_scenario[scenario]
            if not remaining:
                continue
            candidate = remaining.pop(0)
            candidate_id = str(candidate["id"])
            if candidate_id in seen:
                continue
            ordered.append((candidate, scenario))
            seen.add(candidate_id)
            if len(ordered) >= count:
                break

    return ordered[:count]


def _clean_cv_text(text: str, full_name: str, email: str, index: int) -> str:
    lines = text.splitlines()
    cleaned: list[str] = []
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("HỒ SƠ ỨNG VIÊN HƯ CẤU") or stripped.startswith("Mã hồ sơ:"):
            continue
        if stripped.startswith("Họ tên:"):
            cleaned.append(f"Họ tên: {full_name}")
        elif stripped.startswith("Email:"):
            cleaned.append(f"Email: {email}")
        elif stripped.startswith("Điện thoại:"):
            cleaned.append(f"Điện thoại: 09{index:08d}")
        else:
            cleaned.append(line)
    return "\n".join(cleaned).strip()


def _split_sections(text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    title = "HỒ SƠ ỨNG VIÊN"
    body: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            if body and body[-1] != "":
                body.append("")
            continue
        is_heading = (
            len(line) <= 70
            and line == line.upper()
            and any(character.isalpha() for character in line)
            and not line.startswith(("EMAIL:", "HỌ TÊN:", "ĐIỆN THOẠI:"))
        )
        if is_heading and body:
            sections.append((title, body))
            title, body = line, []
        elif is_heading:
            title = line
        else:
            body.append(line)
    if body:
        sections.append((title, body))
    return sections


def _set_cell_shading(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shade)


def _set_columns(section: Any, count: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    section_properties = section._sectPr
    existing = section_properties.xpath("./w:cols")
    columns = existing[0] if existing else OxmlElement("w:cols")
    columns.set(qn("w:num"), str(count))
    if not existing:
        section_properties.append(columns)


def _add_line(document: Any, line: str, compact: bool = False) -> None:
    paragraph = document.add_paragraph(line)
    paragraph.paragraph_format.space_after = 0 if compact else 3
    paragraph.paragraph_format.line_spacing = 1.0 if compact else 1.08


def _write_docx(
    target: Path,
    full_name: str,
    text: str,
    layout_index: int,
) -> str:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    sections = _split_sections(text)
    variant = layout_index
    layout_names = {
        0: "một cột tiêu chuẩn",
        1: "tiêu đề màu",
        2: "bảng thông tin",
        3: "hai cột bằng bảng",
        4: "kỹ năng dạng bảng",
        5: "timeline dạng bảng",
        6: "hai cột Word",
        7: "trang đầu ngang",
        8: "nhiều bảng có ô tiêu đề gộp",
        15: "các khối nội dung có thanh tiêu đề",
        16: "sidebar hồ sơ và nội dung chính",
        19: "bảng nhãn dọc theo từng mục",
    }

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if variant in {0, 1, 4} else WD_ALIGN_PARAGRAPH.LEFT
    name_run = heading.add_run(full_name.upper())
    name_run.bold = True
    name_run.font.size = Pt(19 if variant != 6 else 16)
    name_run.font.color.rgb = RGBColor(22, 78, 99) if variant in {1, 4, 6} else RGBColor(15, 23, 42)

    if variant == 7:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    if variant == 15:
        for section_index, (title, lines) in enumerate(sections):
            table = document.add_table(rows=2, cols=1)
            table.autofit = True
            header, content = table.rows[0].cells[0], table.rows[1].cells[0]
            header.text = title
            _set_cell_shading(header, "0F766E" if section_index % 2 == 0 else "155E75")
            header.paragraphs[0].runs[0].bold = True
            header.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            content.text = ""
            for line in [item for item in lines if item]:
                paragraph = content.add_paragraph(style=None)
                paragraph.add_run(f"• {line}")
                paragraph.paragraph_format.space_after = Pt(2)
            document.add_paragraph().paragraph_format.space_after = Pt(1)
    elif variant == 16:
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        sidebar, main = table.rows[0].cells
        sidebar.width, main.width = Cm(6.2), Cm(11.2)
        _set_cell_shading(sidebar, "ECFEFF")
        sidebar_sections = sections[:3]
        main_sections = sections[3:]
        for title, lines in sidebar_sections:
            paragraph = sidebar.add_paragraph(title)
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.color.rgb = RGBColor(14, 116, 144)
            for line in [item for item in lines if item]:
                sidebar.add_paragraph(line)
        for title, lines in main_sections:
            paragraph = main.add_paragraph(title)
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.color.rgb = RGBColor(15, 23, 42)
            for line in [item for item in lines if item]:
                main.add_paragraph(line)
    elif variant == 19:
        section.left_margin = Cm(1.2)
        section.right_margin = Cm(1.2)
        for section_index, (title, lines) in enumerate(sections):
            table = document.add_table(rows=1, cols=2)
            table.autofit = False
            label, content = table.rows[0].cells
            label.width, content.width = Cm(4.1), Cm(14.2)
            label.text = title
            label.paragraphs[0].runs[0].bold = True
            _set_cell_shading(label, "FEF3C7" if section_index % 2 == 0 else "E0E7FF")
            content.text = "\n".join(item for item in lines if item)
            document.add_paragraph().paragraph_format.space_after = Pt(1)
    elif variant == 3:
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        left, right = table.rows[0].cells
        left.width, right.width = Cm(6), Cm(11)
        midpoint = max(1, len(sections) // 3)
        for title, lines in sections[:midpoint]:
            paragraph = left.add_paragraph(title)
            paragraph.runs[0].bold = True
            _set_cell_shading(left, "E0F2FE")
            for line in lines:
                if line:
                    left.add_paragraph(line)
        for title, lines in sections[midpoint:]:
            paragraph = right.add_paragraph(title)
            paragraph.runs[0].bold = True
            for line in lines:
                if line:
                    right.add_paragraph(line)
    elif variant in {4, 5}:
        for title, lines in sections:
            paragraph = document.add_heading(title, level=1)
            paragraph.runs[0].font.size = Pt(12)
            if (variant == 4 and "KỸ NĂNG" in title) or (variant == 5 and "KINH NGHIỆM" in title):
                table = document.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for line_number, line in enumerate([item for item in lines if item]):
                    cells = table.add_row().cells
                    cells[0].text = str(line_number + 1)
                    cells[1].text = line
                    if line_number % 2 == 0:
                        _set_cell_shading(cells[1], "F8FAFC")
            else:
                for line in lines:
                    if line:
                        _add_line(document, line)
    elif variant == 8:
        for section_index, (title, lines) in enumerate(sections):
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            header = table.rows[0].cells[0].merge(table.rows[0].cells[1])
            header.text = title
            _set_cell_shading(header, "DBEAFE" if section_index % 2 == 0 else "E2E8F0")
            header.paragraphs[0].runs[0].bold = True
            content_lines = [line for line in lines if line]
            for line_index in range(0, len(content_lines), 2):
                cells = table.add_row().cells
                cells[0].text = content_lines[line_index]
                cells[1].text = content_lines[line_index + 1] if line_index + 1 < len(content_lines) else ""
            document.add_paragraph().paragraph_format.space_after = Pt(2)
    else:
        if variant == 6:
            _set_columns(section, 2)
        if variant == 2 and sections:
            contact_lines = [line for line in sections[0][1] if ":" in line][:6]
            if contact_lines:
                table = document.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for line in contact_lines:
                    key, value = line.split(":", 1)
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = key, value.strip()
                    _set_cell_shading(cells[0], "E2E8F0")
        for title, lines in sections:
            paragraph = document.add_heading(title, level=1)
            paragraph.runs[0].font.size = Pt(12 if variant != 1 else 13)
            if variant == 1:
                paragraph.runs[0].font.color.rgb = RGBColor(3, 105, 161)
            for line in lines:
                if line:
                    _add_line(document, line, compact=variant == 6)
        if variant == 7:
            new_section = document.add_section(WD_SECTION.NEW_PAGE)
            new_section.orientation = WD_ORIENT.PORTRAIT
            new_section.page_width, new_section.page_height = section.page_height, section.page_width

    core = document.core_properties
    core.title = f"CV {full_name}"
    core.subject = "Hồ sơ kiểm thử end-to-end tuyển dụng"
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return layout_names[variant]


def _font_paths() -> tuple[Path, Path]:
    candidates = (
        (Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
    )
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            return regular, bold
    raise FileNotFoundError("Không tìm thấy font Unicode Arial hoặc DejaVu Sans để sinh PDF kiểm thử.")


def _wrapped_pdf_lines(text: str, width: int) -> list[tuple[str, bool]]:
    wrapped: list[tuple[str, bool]] = []
    for raw in text.splitlines():
        value = raw.strip()
        if not value:
            wrapped.append(("", False))
            continue
        heading = (
            len(value) <= 70
            and value == value.upper()
            and any(character.isalpha() for character in value)
        )
        chunks = textwrap.wrap(value, width=width, break_long_words=False) or [""]
        wrapped.extend((chunk, heading) for chunk in chunks)
    return wrapped


def _write_text_pdf(target: Path, text: str, variant: int) -> str:
    from fpdf import FPDF

    regular, bold = _font_paths()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_font("CvArial", "", str(regular))
    pdf.add_font("CvArial", "B", str(bold))
    if variant == 21:
        sections = _split_sections(text)
        pdf.add_page(orientation="L")
        column_width = 132
        x_positions = (12, 151)
        column_index = 0
        y = 12
        for title, section_lines in sections:
            estimated_height = 8 + max(1, len(section_lines)) * 4.2
            if y + estimated_height > 195:
                column_index += 1
                if column_index >= len(x_positions):
                    pdf.add_page(orientation="L")
                    column_index = 0
                y = 12
            x = x_positions[column_index]
            pdf.set_xy(x, y)
            pdf.set_font("CvArial", "B", 9.4)
            pdf.set_fill_color(219, 234, 254)
            pdf.multi_cell(column_width, 5.6, title, border=1, fill=True)
            y = pdf.get_y()
            pdf.set_font("CvArial", "", 7.5)
            for line in [item for item in section_lines if item]:
                pdf.set_xy(x, y)
                pdf.multi_cell(column_width, 4.2, line, border="B")
                y = pdf.get_y()
            y += 2
        description = "PDF text ngang hai cột theo khối nội dung"
    elif variant == 20:
        lines = _wrapped_pdf_lines(text, 50)
        lines_per_column = 54
        lines_per_page = lines_per_column * 2
        for page_start in range(0, len(lines), lines_per_page):
            pdf.add_page()
            page_lines = lines[page_start:page_start + lines_per_page]
            for column_index in range(2):
                column_lines = page_lines[
                    column_index * lines_per_column:(column_index + 1) * lines_per_column
                ]
                x = 10 + column_index * 100
                y = 12
                for value, heading in column_lines:
                    pdf.set_xy(x, y)
                    pdf.set_font("CvArial", "B" if heading else "", 9.1 if heading else 7.6)
                    if heading:
                        pdf.set_fill_color(241, 245, 249)
                        pdf.cell(92, 5.4, value, border=1, fill=True)
                        y += 5.8
                    else:
                        pdf.multi_cell(92, 4.4, value, border="B")
                        y = pdf.get_y() + 0.3
        description = "PDF text hai cột có khung mục mở rộng"
    elif variant == 17:
        lines = _wrapped_pdf_lines(text, 43)
        lines_per_column = 40
        lines_per_page = lines_per_column * 3
        for page_start in range(0, len(lines), lines_per_page):
            pdf.add_page(orientation="L")
            page_lines = lines[page_start:page_start + lines_per_page]
            for column_index in range(3):
                column_lines = page_lines[
                    column_index * lines_per_column:(column_index + 1) * lines_per_column
                ]
                x = 10 + column_index * 95
                y = 12
                for value, heading in column_lines:
                    pdf.set_xy(x, y)
                    pdf.set_font("CvArial", "B" if heading else "", 9 if heading else 7.4)
                    if heading:
                        pdf.set_fill_color(224, 242, 254)
                        pdf.cell(88, 5.1, value, border=1, fill=True)
                    else:
                        pdf.multi_cell(88, 4.3, value, border="B")
                    y = pdf.get_y() + 0.4
        description = "PDF text ngang ba cột có đường phân mục"
    elif variant == 9:
        pdf.add_page()
        for value, heading in _wrapped_pdf_lines(text, 105):
            pdf.set_x(pdf.l_margin)
            pdf.set_font("CvArial", "B" if heading else "", 11 if heading else 8.5)
            pdf.multi_cell(0, 4.8 if not heading else 6.2, value)
        description = "PDF text một cột nhiều trang"
    elif variant == 10:
        lines = _wrapped_pdf_lines(text, 54)
        lines_per_column = 57
        lines_per_page = lines_per_column * 2
        for page_start in range(0, len(lines), lines_per_page):
            pdf.add_page()
            page_lines = lines[page_start:page_start + lines_per_page]
            for column_index in range(2):
                column_lines = page_lines[
                    column_index * lines_per_column:(column_index + 1) * lines_per_column
                ]
                x = 10 + column_index * 100
                y = 13
                for value, heading in column_lines:
                    pdf.set_xy(x, y)
                    pdf.set_font("CvArial", "B" if heading else "", 9.2 if heading else 7.7)
                    pdf.cell(92, 4.7, value)
                    y += 4.7
        description = "PDF text hai cột"
    else:
        sections = _split_sections(text)
        pdf.add_page()
        for title, section_lines in sections:
            pdf.set_x(pdf.l_margin)
            pdf.set_font("CvArial", "B", 9.5)
            pdf.set_fill_color(219, 234, 254)
            pdf.multi_cell(0, 6, title, border=1, fill=True)
            pdf.set_font("CvArial", "", 7.7)
            for line in [item for item in section_lines if item]:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 4.4, line, border=1)
            pdf.ln(2)
        description = "PDF text nhiều bảng"
    target.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(target))
    return description


def _write_scan_pdf(target: Path, text: str, variant: int) -> str:
    from fpdf import FPDF
    from PIL import Image, ImageDraw, ImageFont

    regular_path, bold_path = _font_paths()
    two_columns = variant in {13, 18}
    noisy = variant == 14
    width, height = 1654, 2339
    margin_x, start_y = 80, 90
    column_count = 2 if two_columns else 1
    wrap_width = 48 if two_columns else 100
    wrapped = _wrapped_pdf_lines(text, wrap_width)
    lines_per_column = 52
    lines_per_page = lines_per_column * column_count
    pages = max(1, math.ceil(len(wrapped) / lines_per_page))
    regular = ImageFont.truetype(str(regular_path), 24 if not two_columns else 21)
    bold = ImageFont.truetype(str(bold_path), 29 if not two_columns else 25)
    page_paths: list[Path] = []
    try:
        for page_index in range(pages):
            image = Image.new("RGB", (width, height), "white")
            draw = ImageDraw.Draw(image)
            draw.rectangle((45, 45, width - 45, height - 45), outline="#cbd5e1", width=3)
            if variant == 18:
                watermark_font = ImageFont.truetype(str(bold_path), 72)
                draw.text(
                    (250, 1080),
                    "HỒ SƠ ỨNG TUYỂN",
                    font=watermark_font,
                    fill="#edf2f7",
                )
            page_lines = wrapped[page_index * lines_per_page:(page_index + 1) * lines_per_page]
            for column_index in range(column_count):
                current = page_lines[
                    column_index * lines_per_column:(column_index + 1) * lines_per_column
                ]
                x = margin_x + column_index * 790
                y = start_y
                for value, heading in current:
                    if heading:
                        draw.rectangle((x - 8, y - 4, x + (710 if two_columns else 1470), y + 34), fill="#e2e8f0")
                    draw.text((x, y), value, font=bold if heading else regular, fill="#111827")
                    y += 40 if heading else 36
            if noisy:
                for x in range(0, width, 145):
                    draw.line((x, 0, min(width, x + 440), height), fill="#f1f5f9", width=2)
                image = image.rotate(2.0, expand=False, fillcolor="white")
            page_path = target.parent / f".{target.stem}-page-{page_index + 1}.jpg"
            image.save(page_path, format="JPEG", quality=88 if noisy else 94)
            page_paths.append(page_path)

        pdf = FPDF()
        for page_path in page_paths:
            pdf.add_page()
            pdf.image(str(page_path), x=0, y=0, w=210, h=297)
        target.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(target))
    finally:
        for page_path in page_paths:
            page_path.unlink(missing_ok=True)
    return {
        12: "PDF scan rõ nét nhiều trang",
        13: "PDF scan hai cột nhiều trang",
        14: "PDF scan lệch và có nhiễu",
        18: "PDF scan hai cột tương phản thấp có watermark",
    }[variant]


def _write_candidate_document(
    output_dir: Path,
    file_stem: str,
    full_name: str,
    text: str,
    layout_index: int,
    variant_override: int | None = None,
) -> tuple[Path, str]:
    variant = variant_override if variant_override is not None else LAYOUT_VARIANT_ORDER[layout_index % len(LAYOUT_VARIANT_ORDER)]
    if variant <= 8 or variant in {15, 16, 19}:
        target = output_dir / f"{file_stem}.docx"
        return target, _write_docx(target, full_name, text, variant)
    target = output_dir / f"{file_stem}.pdf"
    if variant <= 11 or variant in {17, 20, 21}:
        return target, _write_text_pdf(target, text, variant)
    return target, _write_scan_pdf(target, text, variant)


def build_candidate_fixtures(
    workspace: Path,
    output_dir: Path,
    benchmark_job_id: str,
    count: int,
    run_id: str,
    email_domain: str,
) -> list[CandidateFixture]:
    generated = workspace / "Python" / "test_data" / "offline_benchmark" / "generated"
    cvs = _load_json(generated / "cvs.json")
    pairs = _load_json(generated / "pairs.json")
    selected = _select_stratified_pairs(cvs, pairs, benchmark_job_id, count)
    if len(selected) < count:
        raise ValueError(
            f"Benchmark {benchmark_job_id} chỉ có {len(selected)} cặp CV-JD phân tầng, cần {count}."
        )
    expected_scores = _load_expected_scores(
        workspace
        / "Python"
        / "test_data"
        / "offline_benchmark"
        / "results"
        / "matching_results.csv"
    )

    safe_run = re.sub(r"[^a-zA-Z0-9]+", "", run_id).lower()[:24] or "run"
    domain = email_domain.strip().lower()
    if not domain or "@" in domain:
        raise ValueError("E2E_CANDIDATE_EMAIL_DOMAIN phải là tên miền, không chứa @.")

    fixtures: list[CandidateFixture] = []
    for offset, (cv, scenario) in enumerate(selected, start=1):
        full_name = FICTIONAL_NAMES[(offset - 1) % len(FICTIONAL_NAMES)]
        email = f"e2e.{safe_run}.{offset:02d}@{domain}"
        text = _clean_cv_text(str(cv.get("text", "")), full_name, email, offset)
        if len(text.split()) < 900:
            raise ValueError(f"CV {cv['id']} quá ngắn sau khi chuẩn hóa ({len(text.split())} từ).")
        target, layout = _write_candidate_document(
            output_dir,
            f"cv_{offset:02d}_{safe_run}",
            full_name,
            text,
            offset - 1,
        )
        fixtures.append(
            CandidateFixture(
                index=offset,
                source_cv_id=str(cv["id"]),
                full_name=full_name,
                email=email,
                case_type=str(cv.get("case_type", "unknown")),
                match_scenario=scenario,
                expected_offline_score=expected_scores.get(
                    (str(cv["id"]), benchmark_job_id, scenario),
                    -1,
                ),
                document_path=target.resolve(),
                layout=layout,
            )
        )
    return fixtures
