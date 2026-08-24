from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("TAI_LIEU_DEMO_PHAN_BIEN/04_ON_TAP_TOAN_HE_THONG_VA_GIAI_THUAT.md")
TARGET = Path("TAI_LIEU_DEMO_PHAN_BIEN/04_ON_TAP_TOAN_HE_THONG_VA_GIAI_THUAT.docx")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (
        ("Title", 20, "17365D"),
        ("Heading 1", 16, "17365D"),
        ("Heading 2", 14, "1F4E79"),
        ("Heading 3", 12, "2F5597"),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def convert() -> None:
    document = Document()
    configure(document)
    quote_mode = False

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        if not line:
            quote_mode = False
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            add_inline(document.add_paragraph(style="Heading 1"), line[3:])
        elif line.startswith("### "):
            add_inline(document.add_paragraph(style="Heading 2"), line[4:])
        elif re.match(r"^\d+\.\s", line):
            add_inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            add_inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.right_indent = Cm(0.5)
            run = paragraph.add_run(line[2:])
            run.italic = True
            run.font.color.rgb = __import__("docx").shared.RGBColor.from_string("404040")
            quote_mode = True
        else:
            paragraph = document.add_paragraph()
            if quote_mode:
                paragraph.paragraph_format.left_indent = Cm(0.8)
            add_inline(paragraph, line)

    core = document.core_properties
    core.title = "Ôn tập demo và phản biện toàn hệ thống RecruitInsightAI"
    core.subject = "Chức năng, quy trình, giải thuật và câu hỏi phản biện"
    core.author = "Nhóm thực hiện RecruitInsightAI"
    document.save(TARGET)


if __name__ == "__main__":
    convert()
