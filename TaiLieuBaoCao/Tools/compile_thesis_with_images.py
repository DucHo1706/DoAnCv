import os
import re
import sys
import base64
import urllib.request
import urllib.parse
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import zlib

def download_mermaid_image(mermaid_code, output_dir, name):
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, f"{name}.png")
    
    # Kroki compression logic (zlib + urlsafe base64)
    clean_code = mermaid_code.strip()
    try:
        compressed = zlib.compress(clean_code.encode('utf-8'), 9)
        encoded_string = base64.urlsafe_b64encode(compressed).decode('utf-8')
        url = f"https://kroki.io/mermaid/png/{encoded_string}"
        
        headers = {"User-Agent": "Mozilla/5.0"}
        print(f"Rendering Mermaid diagram [{name}] via Kroki API...")
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=40) as response:
            with open(out_path, "wb") as f:
                f.write(response.read())
        print(f"-> Saved: {out_path} ({os.path.getsize(out_path)} bytes)")
        return out_path
    except Exception as e:
        print(f"-> WARNING: Failed to render {name}: {e}. Falling back to text.")
        return None

def build_thesis_docx(md_path, docx_path):
    print(f"Reading markdown from {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Clean XML-incompatible control characters
    content = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', content)

    doc = docx.Document()
    
    # Page Margins (Left 3cm, Top/Bottom/Right 2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.79)

    # Base Style Config
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    def add_chapter_heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(71, 85, 105)
        return p

    def add_heading_4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(71, 85, 105)
        return p

    def add_body_paragraph(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Simple inline formatting parser (bold/italic)
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)
        
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
        return p

    def add_bullet_point(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)
                
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.1)
        cell = tbl.cell(0, 0)
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
        
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:left w:val="single" w:sz="16" w:space="0" w:color="2563EB"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'</w:tcBorders>'
        )
        cell._tc.get_or_add_tcPr().append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        run = p.add_run(code_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(71, 85, 105)

    def add_table_data(headers, rows):
        tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
        tbl.autofit = True
        
        # Style table headers
        hdr_cells = tbl.rows[0].cells
        for idx, header in enumerate(headers):
            hdr_cells[idx].text = header
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.runs[0]
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            
            shading_xml = f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'
            hdr_cells[idx]._tc.get_or_add_tcPr().append(parse_xml(shading_xml))
            
        # Add table data
        for r_idx, row in enumerate(rows):
            row_cells = tbl.rows[r_idx+1].cells
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = val
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                if p.runs:
                    run = p.runs[0]
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    
        # Apply borders
        for row in tbl.rows:
            for cell in row.cells:
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'</w:tcBorders>'
                )
                cell._tc.get_or_add_tcPr().append(tcBorders)

    def add_image_block(image_path, caption):
        # Insert image
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        try:
            run = p.add_run()
            run.add_picture(image_path, width=Inches(5.8))
        except Exception as e:
            p.add_run(f"[Không thể tải ảnh: {os.path.basename(image_path)}. Lỗi: {e}]")
            
        # Insert caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.italic = True
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139) # Slate-500

    # Split lines
    lines = [line.strip() for line in content.split('\n')]
    
    image_dir = r"d:\KhoaLuan\TaiLieuBaoCao\Documents\images"
    uml_img_dir = os.path.join(image_dir, "uml")
    
    # Custom mapping for Mermaid diagram names and captions
    mermaid_diagram_specs = [
        {"name": "architecture_diagram", "caption": "Hình 3.1: Sơ đồ kiến trúc tổng thể hệ thống tuyển dụng AI"},
        {"name": "usecase_diagram", "caption": "Hình 3.2: Sơ đồ Use Case tổng quát hệ thống"},
        {"name": "class_diagram", "caption": "Hình 3.3: Sơ đồ lớp thực thể hệ thống (Class Diagram)"},
        {"name": "sequence_diagram_1", "caption": "Hình 3.4: Sơ đồ tuần tự quy trình Ứng viên nộp đơn & Đánh giá AI"},
        {"name": "sequence_diagram_2", "caption": "Hình 3.5: Sơ đồ tuần tự quy trình Đăng tin JD & Cấu hình ASK"},
        {"name": "sequence_diagram_3", "caption": "Hình 3.6: Sơ đồ tuần tự quy trình Xem Bảng xếp hạng CV & Truy vấn bản đồ Apriori"},
        {"name": "activity_diagram_1", "caption": "Hình 3.7: Sơ đồ hoạt động chạy ngầm huấn luyện lại giải thuật"},
        {"name": "activity_diagram_2", "caption": "Hình 3.8: Sơ đồ hoạt động quy trình xử lý phân tích & Chấm điểm CV chi tiết (FastAPI)"},
        {"name": "erd_diagram", "caption": "Hình 3.9: Sơ đồ thực thể mối quan hệ (ERD)"}
    ]
    
    mermaid_counter = 0
    mockup_counter = 10 # UI Mockups start at Hình 3.10
    
    print("Parsing document...")
    idx = 0
    in_code = False
    code_lines = []
    
    in_table = False
    table_headers = []
    table_rows = []

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        
        # 1. Code blocks (Mermaid or code)
        if stripped.startswith('```'):
            if in_code:
                code_text = '\n'.join(code_lines)
                if in_code_type == "mermaid" and mermaid_counter < len(mermaid_diagram_specs):
                    # It's a Mermaid diagram, render it!
                    spec = mermaid_diagram_specs[mermaid_counter]
                    img_path = download_mermaid_image(code_text, uml_img_dir, spec["name"])
                    if img_path and os.path.exists(img_path):
                        add_image_block(img_path, spec["caption"])
                    else:
                        # Fallback to code block
                        add_code_block(code_text)
                    mermaid_counter += 1
                else:
                    add_code_block(code_text)
                in_code = False
                code_lines = []
            else:
                in_code = True
                in_code_type = "mermaid" if "mermaid" in stripped.lower() else "code"
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        # 2. Markdown Image tags parser (e.g. ![caption](file:///d:/...))
        image_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if image_match:
            img_caption = image_match.group(1)
            raw_url = image_match.group(2)
            
            # Convert file URL or relative path to standard local path
            parsed_url = urllib.parse.urlparse(raw_url)
            local_path = urllib.parse.unquote(parsed_url.path)
            if local_path.startswith('/') and os.name == 'nt' and ':' in local_path:
                local_path = local_path.lstrip('/')
                
            # If path is relative, resolve it
            if not os.path.isabs(local_path):
                local_path = os.path.abspath(os.path.join(os.path.dirname(md_path), local_path))
                
            print(f"Parsing mockup image: {local_path}...")
            
            # Format UI Mockup caption
            full_caption = f"Hình 3.{mockup_counter}: {img_caption}"
            add_image_block(local_path, full_caption)
            mockup_counter += 1
            idx += 1
            continue

        # 3. Tables
        if stripped.startswith('|'):
            in_table = True
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                idx += 1
                continue
                
            if not table_headers:
                table_headers = cells
            else:
                table_rows.append(cells)
            idx += 1
            continue
        else:
            if in_table:
                add_table_data(table_headers, table_rows)
                in_table = False
                table_headers = []
                table_rows = []

        if not stripped:
            idx += 1
            continue

        # 4. Heading 1 (Markdown '# ')
        if stripped.startswith('# '):
            heading_val = stripped[2:].upper()
            if "PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG" in heading_val:
                heading_val = "CHƯƠNG 3: PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG"
            add_chapter_heading(heading_val)
            idx += 1
            continue

        # 5. Heading 2, 3, 4 and body paragraphs
        if stripped.startswith('## '):
            add_heading_2(stripped[3:])
        elif stripped.startswith('### '):
            add_heading_3(stripped[4:])
        elif stripped.startswith('#### '):
            add_heading_4(stripped[5:])
        elif stripped.startswith('- '):
            add_bullet_point(stripped[2:])
        else:
            add_body_paragraph(stripped)
            
        idx += 1

    # End-of-file flushes
    if in_table:
        add_table_data(table_headers, table_rows)
    if in_code:
        add_code_block('\n'.join(code_lines))

    # Save output
    try:
        doc.save(docx_path)
        print(f"\nSUCCESS: Fully compiled Word document with diagrams and mockups at: {docx_path}")
    except PermissionError:
        fallback_path = docx_path.replace(".docx", "_Moi.docx")
        try:
            doc.save(fallback_path)
            print(f"\nWARNING: File is locked. Saved to fallback: {fallback_path}")
        except Exception as e2:
            print(f"\nERROR saving fallback docx: {e2}")
    except Exception as e:
        print(f"\nERROR saving docx: {e}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python compile_thesis_with_images.py <input.md> <output.docx>")
        sys.exit(1)
    build_thesis_docx(sys.argv[1], sys.argv[2])
