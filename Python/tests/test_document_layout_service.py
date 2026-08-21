import io
import unittest
from unittest.mock import patch

from docx import Document

from services.document_layout_service import (
    ExtractionCandidate,
    LayoutBlock,
    _order_blocks_for_reading,
    _words_to_blocks,
    choose_best_candidate,
    calculate_text_agreement,
    extract_docx_candidates,
    get_available_ocr_language,
    score_text_quality,
)
import services.document_layout_service as layout_service


class DocumentLayoutServiceTests(unittest.TestCase):
    def test_selects_structured_cv_instead_of_long_garbage(self):
        structured = """NGUYỄN MINH ANH
minhanh@example.com
KINH NGHIỆM LÀM VIỆC
01/2024 - 06/2025
Phát triển REST API bằng ASP.NET Core.
KỸ NĂNG
C#, SQL Server, Docker
HỌC VẤN
Đại học Công nghệ Thông tin
"""
        garbage = "� " * 500
        result = choose_best_candidate([
            ExtractionCandidate(method="garbled", text=garbage),
            ExtractionCandidate(method="structured", text=structured),
        ])
        self.assertEqual(result.method, "structured")
        self.assertGreater(result.quality_score, 45)

    def test_splits_visually_separated_regions_on_same_line(self):
        words = [
            {"text": "KỸ", "x0": 20, "x1": 40, "top": 100, "bottom": 112},
            {"text": "NĂNG", "x0": 45, "x1": 80, "top": 100, "bottom": 112},
            {"text": "KINH", "x0": 360, "x1": 395, "top": 100, "bottom": 112},
            {"text": "NGHIỆM", "x0": 400, "x1": 460, "top": 100, "bottom": 112},
        ]
        blocks = _words_to_blocks(words, 1, "test")
        self.assertEqual([block.text for block in blocks], ["KỸ NĂNG", "KINH NGHIỆM"])

    def test_orders_two_columns_top_to_bottom_instead_of_interleaving_rows(self):
        blocks = []
        for index, text in enumerate(["TRÁI 1", "TRÁI 2", "TRÁI 3"]):
            blocks.append(LayoutBlock(text=text, x0=20, x1=180, top=100 + index * 40, bottom=120 + index * 40))
        for index, text in enumerate(["PHẢI 1", "PHẢI 2", "PHẢI 3"]):
            blocks.append(LayoutBlock(text=text, x0=360, x1=540, top=100 + index * 40, bottom=120 + index * 40))

        ordered, detected = _order_blocks_for_reading(blocks)

        self.assertTrue(detected)
        self.assertEqual(
            [block.text for block in ordered],
            ["TRÁI 1", "TRÁI 2", "TRÁI 3", "PHẢI 1", "PHẢI 2", "PHẢI 3"],
        )

    def test_docx_tables_are_not_dropped(self):
        document = Document()
        document.add_paragraph("HỒ SƠ ỨNG VIÊN")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Kinh nghiệm"
        table.cell(0, 1).text = "01/2024 - 06/2025"
        table.cell(1, 0).text = "Kỹ năng"
        table.cell(1, 1).text = "Docker, Linux"
        stream = io.BytesIO()
        document.save(stream)

        candidates = extract_docx_candidates(stream.getvalue())
        self.assertIn("Kinh nghiệm | 01/2024 - 06/2025", candidates[0].text)
        self.assertIn("Kỹ năng | Docker, Linux", candidates[0].text)

    def test_ocr_language_does_not_silently_use_english_when_vietnamese_pack_is_missing(self):
        layout_service._OCR_LANGUAGE_CACHE = None
        with patch("services.document_layout_service.pytesseract.get_languages", return_value=["eng", "osd"]):
            with self.assertRaisesRegex(RuntimeError, "vie"):
                get_available_ocr_language()
        layout_service._OCR_LANGUAGE_CACHE = None

    def test_ocr_uses_vietnamese_and_english_when_both_are_installed(self):
        layout_service._OCR_LANGUAGE_CACHE = None
        with patch("services.document_layout_service.pytesseract.get_languages", return_value=["vie", "eng", "osd"]):
            self.assertEqual(get_available_ocr_language(), "vie+eng")
        layout_service._OCR_LANGUAGE_CACHE = None

    def test_lost_vietnamese_diacritics_are_marked_as_low_quality(self):
        garbled = (
            "Muc tieu nghe nghiep mong muon ung dung ky nang giao tiep va sang tao noi dung "
            "vao cong viec truyen thong. Kinh nghiem lam viec tai cong ty va tham gia du an. "
            "Hoc van chuyen nganh quan tri kinh doanh. Ky nang tin hoc van phong va lam viec nhom."
        )
        score, warnings = score_text_quality(garbled)
        self.assertLess(score, 45)
        self.assertTrue(any("mất dấu" in warning for warning in warnings))

    def test_agreement_ignores_punctuation_and_layout_order(self):
        left = "KỸ NĂNG: C#, SQL Server. KINH NGHIỆM: Phát triển REST API."
        right = "KINH NGHIỆM\nPhát triển REST API\nKỸ NĂNG\nC# SQL Server"
        self.assertGreater(calculate_text_agreement(left, right), 0.9)

    def test_conflicting_extractors_require_review_even_when_text_looks_valid(self):
        first = " ".join([
            "KINH NGHIỆM Phát triển REST API ASP.NET Core tại công ty phần mềm",
            "KỸ NĂNG C# SQL Server Docker Kubernetes",
            "HỌC VẤN Đại học Công nghệ Thông tin năm 2024",
        ] * 5)
        second = " ".join([
            "KINH NGHIỆM Điều phối sự kiện truyền thông và viết nội dung quảng cáo",
            "KỸ NĂNG MC thiết kế Canva quản lý ngân sách",
            "HỌC VẤN Đại học Khoa học Xã hội năm 2024",
        ] * 5)
        result = choose_best_candidate([
            ExtractionCandidate(method="pypdf2", text=first),
            ExtractionCandidate(method="pdfplumber_plain", text=second),
        ])
        self.assertEqual(result.agreement_kind, "independent_sources")
        self.assertLess(result.agreement_score or 1, 0.62)
        self.assertFalse(result.analysis_safe)


if __name__ == "__main__":
    unittest.main()
