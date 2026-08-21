from __future__ import annotations

import io
import os
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Callable

import PyPDF2
import pytesseract
import numpy as np
from PIL import Image, ImageOps

_OCR_LANGUAGE_CACHE: str | None = None


def _configure_local_tessdata() -> None:
    """Dùng language pack local trên máy dev; Docker/VPS tiếp tục dùng package hệ thống."""
    if os.getenv("TESSDATA_PREFIX"):
        return
    local_tessdata = Path(__file__).resolve().parents[2] / ".local" / "tessdata"
    if (local_tessdata / "vie.traineddata").is_file():
        os.environ["TESSDATA_PREFIX"] = str(local_tessdata)


_configure_local_tessdata()


SECTION_TERMS = {
    "experience": ("kinh nghiệm", "quá trình công tác", "work experience", "employment history", "professional experience"),
    "skills": ("kỹ năng", "năng lực", "technical skills", "core competencies", "technologies"),
    "projects": ("dự án", "projects", "personal projects", "selected projects"),
    "education": ("học vấn", "education", "academic background"),
    "certifications": ("chứng chỉ", "certifications", "certificates"),
    "summary": ("mục tiêu nghề nghiệp", "giới thiệu", "career objective", "professional summary", "profile"),
}


def get_available_ocr_language() -> str:
    """CV Việt/Anh bắt buộc có model tiếng Việt; không âm thầm OCR tiếng Việt bằng model Anh."""
    global _OCR_LANGUAGE_CACHE
    if _OCR_LANGUAGE_CACHE:
        return _OCR_LANGUAGE_CACHE
    try:
        installed = set(pytesseract.get_languages(config=""))
    except Exception:
        installed = set()
    if {"vie", "eng"}.issubset(installed):
        selected = "vie+eng"
    elif "vie" in installed:
        selected = "vie"
    else:
        raise RuntimeError(
            "Thiếu language pack Tesseract 'vie'. Cài tesseract-ocr-vie hoặc cấu hình TESSDATA_PREFIX."
        )
    _OCR_LANGUAGE_CACHE = selected
    return selected


@dataclass
class LayoutBlock:
    text: str
    page: int = 1
    x0: float = 0.0
    top: float = 0.0
    x1: float = 0.0
    bottom: float = 0.0
    block_type: str = "text"
    confidence: float = 1.0
    source: str = "unknown"


@dataclass
class ExtractionCandidate:
    method: str
    text: str
    blocks: list[LayoutBlock] = field(default_factory=list)
    quality_score: float = 0.0
    warnings: list[str] = field(default_factory=list)


@dataclass
class DocumentExtractionResult:
    text: str
    method: str
    quality_score: float
    quality_level: str
    blocks: list[LayoutBlock]
    alternatives: list[dict]
    warnings: list[str]
    agreement_score: float | None = None
    agreement_kind: str = "single_source"
    analysis_safe: bool = False

    def to_dict(self) -> dict:
        value = asdict(self)
        value["blocks"] = [asdict(item) for item in self.blocks]
        return value


def normalize_extracted_text(text: str) -> str:
    lines = []
    previous = None
    repeat_count = 0
    for raw_line in (text or "").replace("\x00", " ").splitlines():
        line = re.sub(r"[ \t]+", " ", raw_line).strip()
        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        # Header/footer lặp liên tiếp không được nhân nhiều lần vào dữ liệu NLP.
        if line == previous:
            repeat_count += 1
            if repeat_count >= 2:
                continue
        else:
            previous = line
            repeat_count = 0
        lines.append(line)
    return "\n".join(lines).strip()


def score_text_quality(text: str) -> tuple[float, list[str]]:
    value = normalize_extracted_text(text)
    warnings = []
    if not value:
        return 0.0, ["Không trích xuất được văn bản."]

    words = re.findall(r"\b\w+\b", value, re.UNICODE)
    alpha = sum(character.isalpha() for character in value)
    printable = sum(character.isprintable() for character in value)
    replacement_count = value.count("�") + value.count("\ufffd")
    mojibake_count = len(re.findall(r"(?:Ã.|Â.|áº|Ä.|�)", value))
    vietnamese_plain_terms = {
        "kinh", "nghiem", "ky", "nang", "hoc", "van", "lam", "viec",
        "muc", "tieu", "trinh", "do", "cong", "ty", "ung", "vien",
        "thong", "tin", "du", "an", "chuyen", "nganh",
    }
    plain_term_hits = len({word.casefold() for word in words} & vietnamese_plain_terms)
    vietnamese_diacritics = len(re.findall(
        r"[ăâđêôơưáàảãạấầẩẫậắằẳẵặéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ]",
        value.casefold(),
    ))
    section_count = sum(
        1 for aliases in SECTION_TERMS.values()
        if any(alias in value.casefold() for alias in aliases)
    )
    contact_bonus = int(bool(re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", value)))
    date_bonus = int(bool(re.search(r"(?:19|20)\d{2}", value)))

    word_score = min(len(words) / 120, 1.0) * 38
    alpha_score = min(alpha / max(len(value), 1) / 0.65, 1.0) * 22
    printable_score = printable / max(len(value), 1) * 10
    structure_score = min(section_count / 4, 1.0) * 20
    signal_score = (contact_bonus + date_bonus) * 5
    mojibake_penalty = min(mojibake_count * 5, 35)
    lost_diacritic_penalty = 28 if plain_term_hits >= 8 and vietnamese_diacritics < 2 else 0
    penalty = min(replacement_count * 3, 30) + mojibake_penalty + lost_diacritic_penalty
    score = max(0.0, min(100.0, word_score + alpha_score + printable_score + structure_score + signal_score - penalty))

    if len(words) < 35:
        warnings.append("Văn bản trích xuất quá ngắn.")
    if section_count < 2:
        warnings.append("Nhận diện được ít mục nội dung CV.")
    if replacement_count or mojibake_count:
        warnings.append("Văn bản có ký tự lỗi mã hóa.")
    if lost_diacritic_penalty:
        warnings.append("Văn bản tiếng Việt có dấu hiệu mất dấu hoặc OCR nhận dạng sai.")
    return round(score, 1), warnings


def _agreement_tokens(text: str) -> list[str]:
    normalized = unicodedata.normalize("NFKC", text or "").casefold()
    return [
        token
        for token in re.findall(r"[^\W_]+", normalized, re.UNICODE)
        if len(token) >= 2 or token.isdigit()
    ]


def calculate_text_agreement(left: str, right: str) -> float:
    """Đo mức trùng nội dung, không phụ thuộc dấu câu hoặc thứ tự khối của CV nhiều cột."""
    left_tokens = Counter(_agreement_tokens(left))
    right_tokens = Counter(_agreement_tokens(right))
    total = sum(left_tokens.values()) + sum(right_tokens.values())
    if total == 0:
        return 0.0
    overlap = sum((left_tokens & right_tokens).values())
    return round((2 * overlap) / total, 3)


def _extraction_family(method: str) -> str:
    value = (method or "").casefold()
    if "gemini" in value or "vision" in value:
        return "vision"
    if "tesseract" in value or "pdf_ocr" in value:
        return "tesseract"
    if "pdfplumber" in value:
        return "pdfplumber"
    if "pypdf" in value:
        return "pypdf"
    if "docx" in value:
        return "docx"
    return value.split("_", 1)[0] or "unknown"


def choose_best_candidate(candidates: list[ExtractionCandidate]) -> DocumentExtractionResult:
    usable = []
    for candidate in candidates:
        candidate.text = normalize_extracted_text(candidate.text)
        candidate.quality_score, auto_warnings = score_text_quality(candidate.text)
        candidate.warnings.extend(item for item in auto_warnings if item not in candidate.warnings)
        if "tesseract" in candidate.method.casefold() and candidate.blocks:
            ocr_confidence = sum(block.confidence for block in candidate.blocks) / len(candidate.blocks)
            if ocr_confidence < 0.65:
                candidate.quality_score = round(max(0.0, candidate.quality_score - (0.65 - ocr_confidence) * 50), 1)
                if "Độ tin cậy ký tự OCR thấp." not in candidate.warnings:
                    candidate.warnings.append("Độ tin cậy ký tự OCR thấp.")
        usable.append(candidate)
    if not usable:
        usable = [ExtractionCandidate(method="none", text="", warnings=["Không có chiến lược trích xuất khả dụng."])]
    comparable = [
        item for item in usable
        if item.quality_score >= 45 and len(_agreement_tokens(item.text)) >= 20
    ]
    agreement_by_candidate: dict[int, tuple[float | None, str]] = {}
    for candidate in usable:
        independent = [
            calculate_text_agreement(candidate.text, peer.text)
            for peer in comparable
            if peer is not candidate and _extraction_family(peer.method) != _extraction_family(candidate.method)
        ]
        variants = [
            calculate_text_agreement(candidate.text, peer.text)
            for peer in comparable
            if peer is not candidate and _extraction_family(peer.method) == _extraction_family(candidate.method)
        ]
        if independent:
            agreement_by_candidate[id(candidate)] = (max(independent), "independent_sources")
        elif variants:
            agreement_by_candidate[id(candidate)] = (max(variants), "same_engine_variants")
        else:
            agreement_by_candidate[id(candidate)] = (None, "single_source")

    def selection_key(item: ExtractionCandidate) -> tuple[float, float, int]:
        agreement, kind = agreement_by_candidate[id(item)]
        consensus_bonus = (agreement or 0.0) * (12 if kind == "independent_sources" else 4)
        return item.quality_score + consensus_bonus, item.quality_score, len(item.text)

    best = max(usable, key=selection_key)
    block_source = best
    if not best.blocks:
        layout_candidates = [
            item for item in usable
            if item.blocks and item.quality_score >= best.quality_score - 8
        ]
        if layout_candidates:
            block_source = max(layout_candidates, key=lambda item: (item.quality_score, len(item.blocks)))
    level = "high" if best.quality_score >= 75 else "partial" if best.quality_score >= 45 else "insufficient"
    agreement_score, agreement_kind = agreement_by_candidate[id(best)]
    requires_review = (
        level != "high"
        or any(
            token in " ".join(best.warnings).casefold()
            for token in ("mã hóa", "mất dấu", "nhận dạng sai", "độ tin cậy ký tự")
        )
        or (agreement_score is not None and agreement_score < 0.62)
        or (
            _extraction_family(best.method) in {"tesseract", "vision"}
            and agreement_kind != "independent_sources"
        )
    )
    return DocumentExtractionResult(
        text=best.text,
        method=best.method if block_source is best else f"{best.method}+{block_source.method}_blocks",
        quality_score=best.quality_score,
        quality_level=level,
        blocks=block_source.blocks,
        alternatives=[
            {
                "method": item.method,
                "quality_score": item.quality_score,
                "length": len(item.text),
                "agreement_score": agreement_by_candidate[id(item)][0],
                "agreement_kind": agreement_by_candidate[id(item)][1],
            }
            for item in sorted(usable, key=lambda item: item.quality_score, reverse=True)
        ],
        warnings=best.warnings + (["Kết quả trích xuất cần được đối chiếu bằng một nguồn đọc khác."] if requires_review else []),
        agreement_score=agreement_score,
        agreement_kind=agreement_kind,
        analysis_safe=level != "insufficient" and not requires_review,
    )


def _words_to_blocks(words: list[dict], page_number: int, source: str) -> list[LayoutBlock]:
    if not words:
        return []
    ordered = sorted(words, key=lambda item: (round(float(item.get("top", 0)) / 4), float(item.get("x0", 0))))
    lines: list[list[dict]] = []
    for word in ordered:
        top = float(word.get("top", 0))
        height = max(1.0, float(word.get("bottom", top)) - top)
        if lines:
            current_top = sum(float(item.get("top", 0)) for item in lines[-1]) / len(lines[-1])
            current_height = sum(
                max(1.0, float(item.get("bottom", 0)) - float(item.get("top", 0)))
                for item in lines[-1]
            ) / len(lines[-1])
            # Dấu tiếng Việt làm bounding box của các từ cùng dòng lệch vài
            # pixel. Ngưỡng cố định 5 px từng tách "Phan Tuấn Kiệt" thành hai
            # dòng rồi đảo thứ tự trên CV hai cột.
            same_line_tolerance = min(18.0, max(6.0, max(height, current_height) * 0.45))
        else:
            current_top = 0.0
            same_line_tolerance = 0.0
        if not lines or abs(top - current_top) > same_line_tolerance:
            lines.append([word])
        else:
            lines[-1].append(word)

    blocks = []
    for line in lines:
        sorted_line = sorted(line, key=lambda item: float(item.get("x0", 0)))
        # Tách một dòng thành nhiều khối nếu có khoảng trắng ngang lớn; hữu ích
        # cho sidebar, timeline và bố cục không dùng bảng thật.
        segments: list[list[dict]] = [[]]
        previous_x1 = None
        median_height = max(8.0, sum(float(item.get("bottom", 0)) - float(item.get("top", 0)) for item in sorted_line) / len(sorted_line))
        for word in sorted_line:
            x0 = float(word.get("x0", 0))
            if previous_x1 is not None and x0 - previous_x1 > median_height * 4:
                segments.append([])
            segments[-1].append(word)
            previous_x1 = float(word.get("x1", x0))
        for segment in segments:
            text = " ".join(str(item.get("text", "")).strip() for item in segment).strip()
            if text:
                blocks.append(LayoutBlock(
                    text=text,
                    page=page_number,
                    x0=min(float(item.get("x0", 0)) for item in segment),
                    top=min(float(item.get("top", 0)) for item in segment),
                    x1=max(float(item.get("x1", 0)) for item in segment),
                    bottom=max(float(item.get("bottom", 0)) for item in segment),
                    source=source,
                ))
    return blocks


def _order_blocks_for_reading(blocks: list[LayoutBlock]) -> tuple[list[LayoutBlock], bool]:
    """Phát hiện hai cột bằng khoảng cách tâm khối và trả thứ tự đọc theo từng cột."""
    if len(blocks) < 6:
        return sorted(blocks, key=lambda block: (block.top, block.x0)), False
    page_left = min(block.x0 for block in blocks)
    page_right = max(block.x1 for block in blocks)
    page_width = max(page_right - page_left, 1.0)
    candidates = [block for block in blocks if (block.x1 - block.x0) < page_width * 0.72]
    if len(candidates) < 6:
        return sorted(blocks, key=lambda block: (block.top, block.x0)), False

    centers = sorted((block.x0 + block.x1) / 2 for block in candidates)
    gaps = [(centers[index + 1] - centers[index], index) for index in range(len(centers) - 1)]
    largest_gap, gap_index = max(gaps, default=(0.0, 0))
    split = (centers[gap_index] + centers[gap_index + 1]) / 2
    if largest_gap < page_width * 0.16:
        return sorted(blocks, key=lambda block: (block.top, block.x0)), False

    left = [block for block in candidates if (block.x0 + block.x1) / 2 < split]
    right = [block for block in candidates if (block.x0 + block.x1) / 2 >= split]
    if len(left) < 3 or len(right) < 3:
        return sorted(blocks, key=lambda block: (block.top, block.x0)), False

    candidate_ids = {id(block) for block in candidates}
    spanning = [block for block in blocks if id(block) not in candidate_ids]
    first_column_top = min(block.top for block in candidates)
    header = [block for block in spanning if block.top <= first_column_top + 12]
    remainder = [block for block in spanning if block not in header]
    ordered = (
        sorted(header, key=lambda block: (block.top, block.x0))
        + sorted(left, key=lambda block: (block.top, block.x0))
        + sorted(right, key=lambda block: (block.top, block.x0))
        + sorted(remainder, key=lambda block: (block.top, block.x0))
    )
    return ordered, True


def _deskew_image(image: Image.Image) -> tuple[Image.Image, float]:
    """Ước lượng góc nghiêng nhỏ bằng projection profile, không cần OpenCV."""
    grayscale = ImageOps.autocontrast(ImageOps.grayscale(image))
    preview = grayscale.copy()
    preview.thumbnail((1200, 1200))
    array = np.asarray(preview)
    threshold = min(220, int(np.percentile(array, 70)))
    binary = array < threshold
    if binary.mean() < 0.005:
        return grayscale, 0.0

    best_angle = 0.0
    best_score = float(np.var(binary.sum(axis=1)))
    for angle in np.arange(-4.0, 4.01, 0.5):
        if abs(float(angle)) < 0.01:
            continue
        rotated = preview.rotate(float(angle), resample=Image.Resampling.BICUBIC, expand=False, fillcolor=255)
        rotated_array = np.asarray(rotated) < threshold
        score = float(np.var(rotated_array.sum(axis=1)))
        if score > best_score:
            best_score = score
            best_angle = float(angle)
    # Không xoay ảnh chỉ vì dao động rất nhỏ của projection score; phép nội suy
    # không cần thiết có thể làm giảm độ sắc nét của ảnh vốn đã thẳng.
    if abs(best_angle) < 1.0:
        return grayscale, 0.0
    return grayscale.rotate(best_angle, resample=Image.Resampling.BICUBIC, expand=True, fillcolor=255), best_angle


def extract_pdf_candidates(file_bytes: bytes) -> list[ExtractionCandidate]:
    candidates = []
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        candidates.append(ExtractionCandidate(method="pypdf2", text=text))
    except Exception as error:
        candidates.append(ExtractionCandidate(method="pypdf2", text="", warnings=[f"PyPDF2: {error}"]))

    try:
        import pdfplumber
        plain_pages, layout_pages, all_blocks = [], [], []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                plain_pages.append(page.extract_text() or "")
                layout_pages.append(page.extract_text(layout=True) or "")
                all_blocks.extend(_words_to_blocks(page.extract_words(use_text_flow=True), page_number, "pdfplumber_words"))
        candidates.append(ExtractionCandidate(method="pdfplumber_plain", text="\n".join(plain_pages), blocks=all_blocks))
        candidates.append(ExtractionCandidate(method="pdfplumber_layout", text="\n".join(layout_pages), blocks=all_blocks))
    except Exception as error:
        candidates.append(ExtractionCandidate(method="pdfplumber", text="", warnings=[f"pdfplumber: {error}"]))
    return candidates


def extract_docx_candidates(file_bytes: bytes) -> list[ExtractionCandidate]:
    candidates = []
    try:
        from docx import Document
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_lines = []
        blocks = [
            LayoutBlock(text=text, page=1, top=float(index), bottom=float(index + 1), block_type="paragraph", source="python_docx")
            for index, text in enumerate(paragraphs)
        ]
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                values = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if any(values):
                    row_text = " | ".join(values)
                    table_lines.append(row_text)
                    ordinal = len(paragraphs) + len(table_lines) - 1
                    blocks.append(LayoutBlock(
                        text=row_text, page=1, top=float(ordinal), bottom=float(ordinal + 1),
                        block_type="table_row", source="python_docx",
                    ))
        candidates.append(ExtractionCandidate(method="python_docx", text="\n".join(paragraphs + table_lines), blocks=blocks))
    except Exception as error:
        candidates.append(ExtractionCandidate(method="python_docx", text="", warnings=[f"python-docx: {error}"]))
    return candidates


def extract_image_candidates(image: Image.Image) -> list[ExtractionCandidate]:
    candidates = []
    ocr_language = get_available_ocr_language()
    prepared_image, deskew_angle = _deskew_image(image)
    for psm in (4, 6, 11):
        try:
            data = pytesseract.image_to_data(
                prepared_image, lang=ocr_language, config=f"--oem 3 --psm {psm}",
                output_type=pytesseract.Output.DICT,
            )
            words = []
            confidences = []
            for index, raw_text in enumerate(data.get("text", [])):
                text = str(raw_text).strip()
                try:
                    confidence = float(data["conf"][index])
                except (ValueError, TypeError, KeyError):
                    confidence = -1
                if text and confidence >= 0:
                    left, top = float(data["left"][index]), float(data["top"][index])
                    width, height = float(data["width"][index]), float(data["height"][index])
                    words.append({"text": text, "x0": left, "top": top, "x1": left + width, "bottom": top + height})
                    confidences.append(confidence / 100)
            blocks = _words_to_blocks(words, 1, f"tesseract_psm_{psm}")
            blocks, detected_columns = _order_blocks_for_reading(blocks)
            for block in blocks:
                block.confidence = round(sum(confidences) / len(confidences), 2) if confidences else 0.0
            candidates.append(ExtractionCandidate(
                method=(
                    f"tesseract_psm_{psm}"
                    + (f"_deskew_{deskew_angle:+.1f}" if deskew_angle else "")
                    + ("_columns" if detected_columns else "")
                ),
                text="\n".join(block.text for block in blocks),
                blocks=blocks,
            ))
        except Exception as error:
            candidates.append(ExtractionCandidate(method=f"tesseract_psm_{psm}", text="", warnings=[str(error)]))
    return candidates
